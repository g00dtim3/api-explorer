"""Formats dérivés : fichier plat (identique au code existant) et Excel (mêmes valeurs que le CSV).

Données synthétiques fake_*.
"""

import io
import json

import pytest

pytest.importorskip("streamlit")  # utils dépend de Streamlit
pytest.importorskip("openpyxl")

from bulk_export import derived, downloads  # noqa: E402

fake_records = [
    {"id": "1", "attributes": ["Texture", "Scent"], "attributes positive": ["Texture"], "attributes negative": None,
     "brand": "FAKE", "business indicator": "Sampling Rate", "category": "c", "content origin": "o1",
     "content trad": "Texte; avec \"guillemets\"\net retour", "country": "France", "date": "2024-11-28T15:45:28Z",
     "opinion": "positive", "product": "p", "rating": 5, "source": "s", "subcategory": "x"},
    {"id": "2", "attributes": None, "attributes positive": None, "attributes negative": ["Price"], "brand": "FAKE 🙂",
     "business indicator": None, "category": "c", "content origin": "o2", "content trad": "=SOMME(A1)",
     "country": "Italy", "date": "2023-02-03T00:10:00Z", "opinion": "negative", "product": "p2", "rating": None,
     "source": None, "subcategory": "x"},
    {"id": "3", "attributes": ["Safety", "Composition"], "attributes positive": ["Safety"],
     "attributes negative": ["Composition"], "brand": "FAKE", "business indicator": None, "category": "c",
     "content origin": "o3", "content trad": "t3", "country": "France", "date": "2024-01-31T23:59:59Z",
     "opinion": "neutral", "product": "p", "rating": 3, "source": "s", "subcategory": "x", "champ rare": "z"},
    {"id": "4", "attributes": [], "attributes positive": None, "attributes negative": None, "brand": "FAKE",
     "business indicator": "Autre", "category": "c", "content origin": "o4", "content trad": "t4",
     "country": "France", "date": "2024-05-01T08:00:00Z", "opinion": "positive", "product": "p", "rating": 4,
     "source": "s", "subcategory": "x", "imbriqué": {"a": 1, "b": {"c": "d"}}},
]


def mentions(records=fake_records):
    return lambda: iter([(i + 1, f"h{i}", json.dumps(r, ensure_ascii=False)) for i, r in enumerate(records)])


def legacy_flat(records) -> str:
    """Chaîne existante des modules 2/3 : json_normalize + str() des listes + postprocess_reviews + to_csv."""
    import pandas as pd

    from utils import postprocess_reviews

    df = pd.json_normalize(records)
    conv = lambda x: str(x) if isinstance(x, (dict, list)) else x  # noqa: E731
    df = df.map(conv) if hasattr(df, "map") else df.applymap(conv)
    return postprocess_reviews(df.copy()).to_csv(index=False, sep=";")


@pytest.mark.parametrize("chunk", [1, 2, 3, 1000])
def test_plat_par_paquets_identique_au_code_existant(tmp_path, chunk):
    dest = tmp_path / "plat.csv"
    rows = derived.write_plat(mentions(), dest, chunk=chunk)
    assert rows == len(fake_records)
    assert dest.read_bytes().startswith(b"\xef\xbb\xbf")
    assert dest.read_text(encoding="utf-8-sig") == legacy_flat(fake_records)


def test_plat_note_entiere_partout_si_aucune_manquante(tmp_path):
    records = [dict(r, rating=5) for r in fake_records]
    dest = tmp_path / "plat.csv"
    derived.write_plat(mentions(records), dest, chunk=1)
    assert dest.read_text(encoding="utf-8-sig") == legacy_flat(records)


def test_excel_memes_valeurs_que_le_csv(tmp_path):
    from openpyxl import load_workbook

    csv_file = tmp_path / "c.csv"
    _, header = downloads.write_csv(mentions(), csv_file)
    xlsx = tmp_path / "e.xlsx"
    stats = derived.write_excel(mentions(), xlsx, header, len(fake_records))
    assert stats.rows == 4 and stats.truncated_cells == 0 and stats.removed_control_chars == 0
    ws = load_workbook(xlsx, read_only=True)["Data"]
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == header
    for record, row in zip(fake_records, rows[1:]):
        row = list(row) + [None] * (len(header) - len(row))  # openpyxl omet les cellules vides en fin de ligne
        expected = [downloads.csv_cell(record.get(c)) for c in header]
        assert ["" if v is None else (v if isinstance(v, str) else downloads.csv_cell(v)) for v in row] == expected
    assert rows[2][header.index("content trad")] == "=SOMME(A1)"  # texte, pas une formule
    assert rows[1][header.index("rating")] == 5                     # nombre gardé en nombre


def test_excel_cellule_trop_longue_et_caracteres_interdits(tmp_path):
    from openpyxl import load_workbook

    long_text = "x" * 40_000
    records = [{"id": "1", "content trad": long_text, "brand": "FAKE\x01\x02"}]
    xlsx = tmp_path / "e.xlsx"
    stats = derived.write_excel(mentions(records), xlsx, ["id", "content trad", "brand"], 1)
    assert stats.truncated_cells == 1 and stats.removed_control_chars == 2
    row = list(load_workbook(xlsx, read_only=True)["Data"].iter_rows(values_only=True))[1]
    assert len(row[1]) == derived.EXCEL_MAX_CELL and row[2] == "FAKE"


def test_excel_trop_de_lignes_refuse(tmp_path):
    with pytest.raises(derived.ExcelTooLarge):
        derived.write_excel(mentions([]), tmp_path / "e.xlsx", ["id"], derived.EXCEL_MAX_ROWS)
