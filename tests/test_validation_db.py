"""Validation de bout en bout : collecte (faux serveur) → fichiers (stockage local) → contrôles → statut.

Nécessite BULK_EXPORT_TEST_DATABASE_URL. Données synthétiques fake_*.
"""

import gzip
import json
import os
import zipfile

import pytest

from bulk_export import db, finalize, state_store as ss
from bulk_export.api_adapter import RatingsReviewsAdapter
from bulk_export.api_http import RatingsReviewsHttpClient
from bulk_export.models import ExportConfig
from bulk_export.retry import RetryPolicy
from bulk_export.runner import ExportRunner
from bulk_export.settings import load_settings
from bulk_export.storage import LocalStorage, StorageError
from tests.fake_api import FAKE_TOKEN, FakePagedServer
from tests.test_engine_db import fake_doc, fake_pages, make_engine

pytestmark = pytest.mark.db
TEST_SCHEMA = "bulk_export_test"
fake_config = ExportConfig(name="fake validation", filters={"brand": "FAKE_BRAND"}, rows=3, batch_size=10_000)


@pytest.fixture
def connect(pg_conn):
    settings = load_settings({"database_url": os.environ["BULK_EXPORT_TEST_DATABASE_URL"],
                              "db_schema": TEST_SCHEMA})
    return lambda: db.connect(settings, attempts=1)


def adapter_for(server, config=fake_config):
    http = RatingsReviewsHttpClient(FAKE_TOKEN, session=server)
    return RatingsReviewsAdapter(http, config.filters, rows=config.rows, retry_policy=RetryPolicy(max_attempts=1),
                                 sleep=lambda s: None)


def collect(pg_conn, connect, pages, config=fake_config, nb_docs=None):
    export_id = ss.create_export(pg_conn, config)
    result = make_engine(connect, FakePagedServer(pages, nb_docs=nb_docs)).run(export_id)
    assert result.outcome == "FINISHED"
    return export_id


def run_finalize(pg_conn, export_id, storage, server):
    return finalize.finalize(pg_conn, export_id, storage, adapter_for(server))


def check(row, name):
    return next(c for c in row["validation_json"]["checks"] if c["name"] == name)


# ─── Cas nominal ──────────────────────────────────────────────────────────────

def test_export_coherent_completed_fichiers_et_purge(pg_conn, connect, tmp_path):
    pages = fake_pages(total=8)
    export_id = collect(pg_conn, connect, pages)
    store = LocalStorage(tmp_path / "store")
    result = run_finalize(pg_conn, export_id, store, FakePagedServer(pages))
    row = ss.get_export(pg_conn, export_id)
    assert result.passed and row["status"] == "COMPLETED" and row["validation_json"]["passed"] is True
    kinds = sorted(f["kind"] for f in ss.list_files(pg_conn, export_id))
    assert kinds == ["batch_jsonl", "csv_complete", "csv_plat", "manifest", "xlsx", "zip"]
    # Formats dérivés : Excel lisible avec une ligne par mention ; fichier plat identique au code existant.
    from openpyxl import load_workbook

    from tests.test_derived import legacy_flat
    rows = list(load_workbook(store.root / f"{export_id}/output/export_complete.xlsx", read_only=True)["Data"].iter_rows())
    assert len(rows) == 9
    plat = (store.root / f"{export_id}/output/export_plat.csv").read_text(encoding="utf-8-sig")
    assert plat == legacy_flat([d for p in pages for d in p])
    names = {c["name"] for c in row["validation_json"]["checks"]}
    assert {"Intégrité de l'Excel", "Cohérence du fichier plat"} <= names
    assert row["validation_json"]["derived_formats"]["xlsx"]["truncated_cells"] == 0
    # Les fichiers relus redonnent exactement les mentions reçues.
    batch = store.root / f"{export_id}/batches/batch_0001.jsonl.gz"
    lines = gzip.open(batch, "rb").read().decode().splitlines()
    assert [json.loads(line) for line in lines] == [d for p in pages for d in p]
    manifest = json.loads((store.root / f"{export_id}/output/manifest.json").read_text())
    assert manifest["status"] == "COMPLETED" and manifest["volumes"]["unique"] == 8
    assert FAKE_TOKEN not in json.dumps(manifest)
    with zipfile.ZipFile(store.root / f"{export_id}/output/export_batches.zip") as zf:
        assert {"manifest.json", "batches/batch_0001.jsonl", "export_complete.csv"} <= set(zf.namelist())
    # Rétention : mentions effacées de la base, mesures gardées.
    assert pg_conn.execute("SELECT count(*) FROM raw_mentions WHERE export_id = %s", (export_id,)).fetchone()[0] == 0
    assert row["raw_purged_at"] is not None and row["validation_json"]["stats"]["mentions"] == 8


def test_regle_a_doublon_trace_completed(pg_conn, connect, tmp_path):
    p1 = [fake_doc(1), fake_doc(2)]
    p2 = [{**fake_doc(1), "id": "999"}, fake_doc(3)]
    export_id = collect(pg_conn, connect, [p1, p2])  # 4 annoncées, 3 uniques + 1 doublon tracé
    run_finalize(pg_conn, export_id, LocalStorage(tmp_path / "s"), FakePagedServer([p1, p2]))
    row = ss.get_export(pg_conn, export_id)
    assert row["status"] == "COMPLETED"
    manifest = json.loads((tmp_path / "s" / export_id / "output/manifest.json").read_text())
    assert manifest["duplicates"][0]["kept_api_id"] == "1" and manifest["duplicates"][0]["dropped_api_id"] == "999"


# ─── Cas d'échec ──────────────────────────────────────────────────────────────

def test_mention_manquante_incomplete(pg_conn, connect, tmp_path):
    pages = fake_pages(total=8)
    export_id = collect(pg_conn, connect, pages, nb_docs=9)  # l'API annonce 9, en donne 8
    run_finalize(pg_conn, export_id, LocalStorage(tmp_path / "s"), FakePagedServer(pages, nb_docs=9))
    row = ss.get_export(pg_conn, export_id)
    assert row["status"] == "INCOMPLETE" and check(row, "Complétude (règle A)")["ok"] is False
    assert pg_conn.execute("SELECT count(*) FROM raw_mentions WHERE export_id = %s", (export_id,)).fetchone()[0] == 8


def test_jeu_de_donnees_modifie_pendant_l_export_incomplete(pg_conn, connect, tmp_path):
    pages = fake_pages(total=8)
    export_id = collect(pg_conn, connect, pages)
    run_finalize(pg_conn, export_id, LocalStorage(tmp_path / "s"), FakePagedServer(pages, nb_docs=10))
    row = ss.get_export(pg_conn, export_id)
    assert row["status"] == "INCOMPLETE" and check(row, "Données stables pendant l'export")["ok"] is False


def test_mention_rejetee_incomplete(pg_conn, connect, tmp_path):
    pages = [[fake_doc(1), fake_doc(2, date="")]]
    export_id = collect(pg_conn, connect, pages)
    run_finalize(pg_conn, export_id, LocalStorage(tmp_path / "s"), FakePagedServer(pages))
    row = ss.get_export(pg_conn, export_id)
    assert row["status"] == "INCOMPLETE" and check(row, "Aucune mention rejetée")["ok"] is False


class TamperingStorage(LocalStorage):
    """Modifie un lot juste après son envoi (comme si le fichier avait été altéré dans Storage)."""

    def __init__(self, root, mode):
        super().__init__(root)
        self.mode = mode

    def upload_file(self, local_path, remote_path, content_type):
        size = super().upload_file(local_path, remote_path, content_type)
        if remote_path.endswith("batch_0001.jsonl.gz"):
            path = self._path(remote_path)
            lines = gzip.open(path, "rb").read().splitlines(keepends=True)
            if self.mode == "un_caractere":
                lines[0] = lines[0].replace(b"fake", b"fakf", 1)
            elif self.mode == "ligne_dupliquee":
                lines.insert(1, lines[0])
            with gzip.open(path, "wb") as fh:
                fh.writelines(lines)
        return size


@pytest.mark.parametrize("mode", ["un_caractere", "ligne_dupliquee"])
def test_lot_altere_failed(pg_conn, connect, tmp_path, mode):
    pages = fake_pages(total=8)
    export_id = collect(pg_conn, connect, pages)
    run_finalize(pg_conn, export_id, TamperingStorage(tmp_path / "s", mode), FakePagedServer(pages))
    row = ss.get_export(pg_conn, export_id)
    assert row["status"] == "FAILED"
    assert check(row, "Intégrité des lots")["ok"] is False and check(row, "Conservation et unicité")["ok"] is False
    assert row["raw_purged_at"] is None  # rien n'est effacé sans COMPLETED


# ─── Reprise de la validation ─────────────────────────────────────────────────

class BrokenStorage(LocalStorage):
    def upload_file(self, *a, **k):
        raise StorageError("Stockage injoignable (simulé).")


def test_validation_interrompue_puis_relancee(pg_conn, connect, tmp_path):
    pages = fake_pages(total=8)
    export_id = collect(pg_conn, connect, pages)
    first = run_finalize(pg_conn, export_id, BrokenStorage(tmp_path / "b"), FakePagedServer(pages))
    assert not first.passed and ss.get_export(pg_conn, export_id)["status"] == "FAILED"
    again = run_finalize(pg_conn, export_id, LocalStorage(tmp_path / "s"), FakePagedServer(pages))
    assert again.passed and ss.get_export(pg_conn, export_id)["status"] == "COMPLETED"


def test_application_redemarree_pendant_la_validation(pg_conn, connect):
    export_id = collect(pg_conn, connect, fake_pages(total=3))
    ss.acquire_lease(pg_conn, export_id, "app-morte", 60)
    ss.set_status(pg_conn, export_id, "VALIDATING")
    pg_conn.execute("UPDATE exports SET lease_expires_at = now() - interval '1 second'")
    assert ss.recover_orphans(pg_conn) == [export_id]
    assert ss.get_export(pg_conn, export_id)["status"] == "INCOMPLETE"


def test_rien_a_valider_si_collecte_non_terminee(pg_conn, tmp_path):
    export_id = ss.create_export(pg_conn, fake_config)
    result = finalize.finalize(pg_conn, export_id, LocalStorage(tmp_path), adapter_for(FakePagedServer([])))
    assert not result.passed and ss.get_export(pg_conn, export_id)["status"] == "PENDING"


# ─── Rétention, suppression, fichier partiel ──────────────────────────────────

def test_retention_puis_suppression(pg_conn, connect, tmp_path):
    pages = fake_pages(total=3)
    export_id = collect(pg_conn, connect, pages)
    store = LocalStorage(tmp_path / "s")
    run_finalize(pg_conn, export_id, store, FakePagedServer(pages))
    assert finalize.apply_retention(pg_conn, store, retention_days=30) == []
    pg_conn.execute("UPDATE exports SET finished_at = now() - interval '31 days'")
    assert finalize.apply_retention(pg_conn, store, retention_days=30) == [export_id]
    assert not list((store.root / export_id).rglob("*.*"))
    assert ss.get_export(pg_conn, export_id)["files_expired_at"] is not None
    finalize.delete_everywhere(pg_conn, store, export_id)
    with pytest.raises(ss.ExportNotFound):
        ss.get_export(pg_conn, export_id)


def test_suppression_d_un_export_reprenable_demande_confirmation(pg_conn, tmp_path):
    export_id = ss.create_export(pg_conn, fake_config)
    with pytest.raises(ss.StateError, match="confirmation"):
        finalize.delete_everywhere(pg_conn, LocalStorage(tmp_path), export_id)
    ss.acquire_lease(pg_conn, export_id, "quelqu-un", 60)
    with pytest.raises(ss.StateError, match="en cours"):
        finalize.delete_everywhere(pg_conn, LocalStorage(tmp_path), export_id, confirm_resumable=True)
    ss.release_lease(pg_conn, export_id, "quelqu-un")
    finalize.delete_everywhere(pg_conn, LocalStorage(tmp_path), export_id, confirm_resumable=True)


def test_fichier_partiel_etiquete(pg_conn, connect, tmp_path):
    from bulk_export import downloads

    pages = fake_pages(total=8)
    export_id = collect(pg_conn, connect, pages, nb_docs=20)
    store = LocalStorage(tmp_path / "s")
    record = downloads.build_partial_csv(pg_conn, export_id, store, tmp_path)
    assert record.kind == "csv_partial" and record.rows_count == 8 and "DONNEES_PARTIELLES" in record.path
    text = (store.root / record.path).read_text(encoding="utf-8-sig")
    assert text.startswith("DONNÉES PARTIELLES")


# ─── Gestionnaire : collecte puis validation automatique ──────────────────────

def test_runner_collecte_puis_validation_automatique(pg_conn, connect, tmp_path):
    pages = fake_pages(total=8)
    store = LocalStorage(tmp_path / "s")

    def factory(config):
        return adapter_for(FakePagedServer(pages), config)

    runner = ExportRunner(connect, factory, lease_ttl_s=5, heartbeat_s=1, storage_factory=lambda: store)
    export_id = ss.create_export(pg_conn, fake_config)
    assert runner.start(export_id).ok and runner.wait(30)
    row = ss.get_export(pg_conn, export_id)
    assert row["status"] == "COMPLETED" and row["lease_owner"] is None
    # Bouton « Valider » refusé : il n'y a plus rien à valider.
    assert not runner.validate(export_id).ok


class TamperDerived(LocalStorage):
    """Altère l'Excel ou le fichier plat juste après leur envoi."""

    def __init__(self, root, suffix):
        super().__init__(root)
        self.suffix = suffix

    def upload_file(self, local_path, remote_path, content_type):
        size = super().upload_file(local_path, remote_path, content_type)
        if remote_path.endswith(self.suffix):
            with open(self._path(remote_path), "ab") as fh:
                fh.write(b"x")
        return size


@pytest.mark.parametrize("suffix,check_name", [("export_complete.xlsx", "Intégrité de l'Excel"),
                                               ("export_plat.csv", "Intégrité du fichier plat")])
def test_format_derive_altere_failed(pg_conn, connect, tmp_path, suffix, check_name):
    pages = fake_pages(total=5)
    export_id = collect(pg_conn, connect, pages)
    run_finalize(pg_conn, export_id, TamperDerived(tmp_path / "s", suffix), FakePagedServer(pages))
    row = ss.get_export(pg_conn, export_id)
    assert row["status"] == "FAILED" and check(row, check_name)["ok"] is False


def test_formats_derives_non_demandes(pg_conn, connect, tmp_path):
    config = ExportConfig(name="fake sans dérivés", filters={"brand": "FAKE_BRAND"}, rows=3,
                          output_formats=("csv", "jsonl", "zip"))
    pages = fake_pages(total=5)
    export_id = collect(pg_conn, connect, pages, config=config)
    run_finalize(pg_conn, export_id, LocalStorage(tmp_path / "s"), FakePagedServer(pages))
    kinds = {f["kind"] for f in ss.list_files(pg_conn, export_id)}
    assert ss.get_export(pg_conn, export_id)["status"] == "COMPLETED" and not kinds & {"xlsx", "csv_plat"}
